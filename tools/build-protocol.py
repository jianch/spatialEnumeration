"""Build the experiment protocol from the checked production parameter snapshot."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
snapshot = json.loads((ROOT / 'docs/parameters.json').read_text())
C, X = snapshot['config'], snapshot['derived']
T, G, R = C['timing'], C['geometry'], C['training']
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.72)
section.left_margin = section.right_margin = Inches(.75)
section.footer_distance = Inches(.32)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.08
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))): style.element.pPr.remove(border)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Subtitle'].font.size = Pt(10)
for name, size in [('Heading 1', 14), ('Heading 2', 11.5)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = 'Spatial cueing and enumeration experimental design and parameters'
doc.core_properties.subject = 'Behavioural response signal experiment implemented in jsPsych 8'
doc.core_properties.author = ''
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('Page '); run.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text, style=None):
    para = doc.add_paragraph(text, style)
    previous = para._p.getprevious()
    if previous is not None and previous.tag == qn('w:tbl'): para.paragraph_format.space_before = Pt(6)
    return para

def h(text, level=1, new_page=False):
    heading = doc.add_heading(text, level)
    if new_page:
        heading.paragraph_format.page_break_before = True

def table(headers, rows, widths=None):
    widths = widths or [2.30, 4.47]
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths): col.width = Inches(width)
    for i, text in enumerate(headers): tab.rows[0].cells[i].text = text
    repeat = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = tab.add_row().cells
        for i, value in enumerate(values): cells[i].text = str(value)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]: b.set(qn('w:' + key), value)
        borders.append(b)
    tab._tbl.tblPr.append(borders)
    for ri, row in enumerate(tab.rows):
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top','70'), ('bottom','70'), ('left','105'), ('right','105')]:
                item = OxmlElement('w:' + edge); item.set(qn('w:w'), value); item.set(qn('w:type'), 'dxa'); margins.append(item)
            pr.append(margins)
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E5EAF0'); pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.03
                for run in para.runs:
                    run.font.size = Pt(10); run.bold = ri == 0
    return tab

p('Spatial cueing and enumeration\nExperimental design and parameters', 'Title')
p('Full experiment protocol with predictive arrow cues', 'Subtitle')
p(f"jsPsych {C['jsPsychVersion']} | Version {C['experimentVersion']} | 20 October 2026")
p('This experiment tests how preparing attention for a location changes the development of an exact numerical report. A brief dot array is followed by a mask, and a tone at a variable delay instructs the participant to report the number immediately. Valid, invalid, and neutral cues are retained so that benefits of correct preparation and costs of incorrect preparation can be compared separately.')
p('Numerical reports are made with a mouse on an on-screen 0–9 number pad. Arrows appear above a persistent central bull’s-eye plus crosshair. Participants look at the central dot during the dot display; looking behaviour is not recorded. The settings below define the runnable pilot and remain subject to pilot timing and model-recovery evaluation.')
h('Design and allocation')
table(['Factor or allocation', 'Implemented setting'], [
    ('Cue condition', 'Valid, invalid, neutral; within participant'),
    ('Central arrows', '← likely left; → likely right; ↔ either location equally likely; fixed meanings'),
    ('Numerosity', ', '.join(map(str, C['design']['numerosities'])) + ' dots; within participant'),
    ('Signal delay from array onset', ', '.join(map(str, C['design']['signalDelaysMs'])) + ' ms; within participant'),
    ('Target location', 'Left or right, balanced within every cue by numerosity by delay cell in each pilot session'),
    ('Cue proportions', '60% valid, 20% invalid, 20% neutral; directional cues are valid on 75% of directional trials'),
    ('Session structure', f"{X['trials_per_session']} main trials; {X['blocks_per_session']} blocks of {C['design']['blockSize']}; two sessions per participant"),
    ('Total main trials', f"{X['trials_total']} per participant, excluding practice")
])
h('Trials in each numerosity by delay cell', 2)
table(['Cue', 'Per session', 'Across two sessions'], [('Valid',6,12),('Invalid',2,4),('Neutral',2,4),('Total',10,20)], [2.77,2,2])
p('Each session contains 486 valid, 162 invalid, and 162 neutral trials. Each block has ten trials per numerosity. Exact cue proportions apply to cells and the session; individual blocks need not have exactly the same cue proportions.')

h('Trial sequence and timing', new_page=True)
p('The central fixation marker remains visible throughout each timed trial, including the cue display. The dot array appears at one peripheral location and is replaced by masks at both possible locations. The visual sequence continues independently of the response, including when an early tone occurs during the mask.')
table(['Event', 'Timing or rule'], [
    ('Trial start', 'Automatically reset the experiment cursor below fixation, then begin fixation timing'),
    ('Fixation marker', f"Random duration from {T['fixationMinMs']} to {T['fixationMaxMs']} ms"),
    ('Arrow above fixation', f"{T['cueMs']} ms; arrow centre {G['cueOffsetAboveDeg']} degrees above the central dot"),
    ('Post cue fixation', f"{T['postCueMs']} ms; cue onset to array onset is nominally {T['cueMs']+T['postCueMs']} ms"),
    ('Dot array', f"{X['nominal_array_ms']:.1f} ms; {T['arrayFramesAtReference']} frames at {T['referenceHz']} Hz"),
    ('Bilateral masks (masked trials only)', f"{X['nominal_mask_ms']:.0f} ms; {T['maskFramesAtReference']} frames at {T['referenceHz']} Hz"),
    ('Response tone onset', '100, 150, 225, 350, 600, 1000, 1600, 2500, or 3500 ms after array onset'),
    ('Response window', f"0 to {T['responseDeadlineMs']} ms after scheduled tone onset, including both endpoints"),
    ('Late observation', f"A further {T['lateObservationMs']} ms records late presses"),
    ('Selection feedback', f"{T['responseFeedbackMs']} ms highlight; advance only after the target and mask finish"),
    ('Intertrial fixation', f"{T['itiMs']} ms after selection feedback or the unanswered capture window"),
    ('Number pad', 'One horizontal 0–9 row at the reminder centre line; appears at the tone; first left mouse press selects')
])
h('Auditory and display settings', 2)
p(f"The response tone is a {T['toneFrequencyHz']} Hz sine wave lasting {T['toneMs']} ms, with nominal Web Audio gain {T['toneGain']}. Five-millisecond ramps reduce onset and offset clicks. Gain is a software amplitude setting, not a calibrated sound-pressure level. The sound check uses exactly the same pitch, waveform, envelope, and duration as the trial signal.")
p('At the planned refresh rate, the mask begins about 66.7 ms and ends about 166.7 ms after array onset. Therefore, tones at 100 and 150 ms begin while the mask is visible. A response never shortens the target or mask. Selection feedback and the completed mask are followed by the intertrial interval.')
p(f"Visual events use counted animation frames. The full experiment requires a nominal minimum browser frame rate of {T['minimumRefreshHz']} Hz, with {T['pilotRefreshTolerance']*100:.0f}% measurement tolerance (at least {T['minimumRefreshHz']*(1-T['pilotRefreshTolerance']):g} Hz measured). There is no upper-rate limit. The {T['referenceHz']} Hz reference defines the original duration units, not a required display rate. Target and mask durations are converted to whole frames at the measured interval: four and six frames at 60 Hz, eight and twelve at 120 Hz. Other rates are rounded to the nearest whole frame. A critical frame interval exceeding {T['frameGapMultiplier']} times the measured interval is flagged.")
h('Time burden', 2)
p(f"Duration depends on response latency. Omitting every response would take about {X['estimated_main_minutes_without_responses_per_session']:.1f} minutes per session, excluding training, breaks, and interruptions. Prompt selection advances the trial after feedback and the mask, shortening the session. Assess burden in the pilot.")

h('Stimuli and physical setup', new_page=True)
table(['Parameter', 'Value'], [
    ('Array eccentricity', f"{G['eccentricityDeg']} degrees left or right of centre"),
    ('Circular aperture diameter', f"{G['apertureDiameterDeg']} degrees"),
    ('Dot diameter', f"{G['dotDiameterDeg']} degrees"),
    ('Minimum centre spacing', f"{G['minCenterSeparationDeg']} degrees"),
    ('Arrow width and height', f"{G['cueDiameterDeg']} degrees wide and {G['cueDiameterDeg']*0.4:.2f} degrees high"),
    ('Arrow stroke width', f"{G['cueStrokeDeg']} degrees; minimum one CSS pixel"),
    ('Arrow offset above fixation', f"{G['cueOffsetAboveDeg']} degrees to the arrow centre"),
    ('Fixation marker', f"Bull’s-eye plus crosshair: outer diameter {G['fixationDiameterDeg']} degrees; central dot and white band width {G['fixationInnerDiameterDeg']} degrees"),
    ('Mask tile size', f"{G['maskTileDeg']} degrees; grayscale noise within each aperture"),
    ('Colours', f"Background {G['background']}; dots and outlines {G['foreground']}"),
    ('Assumed display', C['displaySetup']['device']),
    ('Brightness setting', '100% ordinary brightness slider; operator assumption, not measured luminance'),
    ('Screen width default', f"{G['defaultScreenWidthCm']} cm; replace with the measured physical width"),
    ('Viewing distance default', f"{G['defaultViewingDistanceCm']} cm; replace with the measured viewing distance")
])
p('Use the built-in MacBook Pro screen at full brightness, with automatic brightness, True Tone and Night Shift off, and power connected. Measure the visible screen width and viewing distance. Use fullscreen and 100% browser zoom. Record the model, display preset, scaling, browser version, and room lighting in the session log. Browser code records these screen settings as assumptions and cannot verify them.')
p('The full experiment supports 60 Hz and faster browser presentation. ProMotion is optional; its adaptive behaviour does not guarantee a particular browser frame rate. Use one display setup consistently across sessions and retain measured frame intervals and per-trial frame records. Validate physical timing before confirmatory data collection. Maximum slider brightness is not the advertised HDR peak or a luminance calibration (Apple Support: support.apple.com/en-au/121554 and support.apple.com/en-us/102297).')
p('At 70 cm, dots are approximately 4.89 mm in diameter, minimum centre spacing is 12.22 mm (7.33 mm edge gap), array and mask diameter is 48.9 mm, and array centres are approximately 61.1 mm left or right under the constant pixels-per-degree mapping. The fixation outer disk is 7.33 mm and its central disk is 2.44 mm. The arrow is nominally 8.55 by 3.42 mm with a 0.86 mm stroke, centred 9.77 mm above fixation; mask tiles are 1.47 mm. Verify the setup page’s 50 mm ruler line and keep viewing distance steady, preferably with a chinrest.')
p('Geometry is informed by Pomè et al. (doi:10.3758/s13414-020-02178-w), Plaisier et al. (2010; doi:10.1007/s00221-010-2319-y), and the fixation target in Thaler et al. (2013), Appendix A (doi:10.1016/j.visres.2012.10.012). The 1 degree spacing and mask tiles are provisional; this layout is not established to eliminate crowding (Chakravarthi and Herbert, 2019; doi:10.1016/j.cognition.2018.12.008).')
p('Coordinates are sampled inside the aperture with minimum spacing. Dot diameter is fixed, so total dot area and aperture density increase with numerosity. Convex hull area, nearest-neighbour spacing, and bounding dimensions are exported as stimulus descriptors; these variables are not independently matched across numerosities.')
h('Response instructions and scoring', new_page=True)
p('Use the existing Continue/Start buttons to begin timed practice and blocks; there is no separate mouse activation notice. Before each trial, the experiment cursor returns below fixation automatically and fixation timing begins. There is no return circle or manual movement step. The reminder reads: Look at the centre dot, only respond after hearing tone. It disappears at target onset. At the tone, immediately left-click your best estimate on the horizontal 0–9 row, centred at the same position as the reminder.')
p(f"The first accepted left mouse press is permanent. The response deadline is {T['responseDeadlineMs']/1000:g} seconds after the tone in timed practice, demonstration, and main trials. Late presses also miss the deadline. Early, right-button, touch, and keyboard inputs cannot select a number. Responses under 100 ms are flagged but retained. Selection highlights for {T['responseFeedbackMs']} ms; the trial advances only after the target and mask finish.")
p(f"The 128 mm response row has 11 by 11 mm buttons, 2 mm gaps, and a top 4 mm below the apertures. Its centre matches the reminder’s 7 mm line box. Number-pad familiarisation uses identical fixed row coordinates, buttons, and gaps as tone practice, full practice, and main trials. Calibrated glyph heights: numerals 4.75 mm; instructions/reminder capitals 3.8 mm; headings 5.5 mm. The pad does not shrink to fit. A drawn arrow cursor resets to {C['mouseCursor']['resetBelowDeg']} degrees below fixation. Browser Pointer Lock hides the system cursor and supplies relative movement; number selection uses virtual viewport CSS pixels. Continue/Start acquires control; it releases for instructions and breaks; Escape releases control and flags an interrupted trial. Exports retain reset coordinates/time, tone/array/pad RT, and pointer position/age at pad onset. Movement after reset is allowed. Main trials give selection highlighting but no correctness feedback.")

h('Training and running a session')
table(['Stage', 'Session order'], [
    ('1. Sound', 'Confirm that the response tone is audible'),
    ('2. Display', 'Measure refresh rate and check calibrated geometry and audio/fullscreen readiness'),
    ('3. Arrow meanings', f"{R['cueTrials']} mouse questions; each of left, right, and neutral once"),
    ('4. Number pad', f"{R['numberPadTrials']} distinct requested numbers; display only Click N"),
    ('5. Tone response', f"{R['toneTrials']} flashing dot trials without masks"),
    ('6. Full practice', f"{R['mixedTrials']} complete dot trials with masks"),
    ('7. Main experiment', 'Proceed to the scheduled main blocks after practice')
])
p(f"All practice sections run once. Feedback lasts {R['feedbackMs']} ms. Wrong, late, missing, and interrupted responses remain in the data; there is no performance threshold or automatic practice repetition. Main blocks follow completion of the practice sequence. Equipment and display checks must still pass. Pilot and demonstration use the same training counts and order.")
p('Practice target exposures are approximately 66.7, 150, and 300 ms: 2/2/1 trials in tone practice and 3/3/2 in full practice. The tone is not planned before target offset. Tone practice clears the target to the background without a mask; full practice presents the bilateral masks. Practice is familiarisation, not a balanced main-design factorial.')
h('Cue meanings', 2)
p('Fixed cue meanings: left arrow predicts left; right arrow predicts right; double-headed arrow gives equal location probability, not simultaneous arrays. The three-question comprehension check presents each meaning once with feedback and no repeat. Answers use mouse buttons.')
p('Arrows share bounds, shaft, stroke, and colour; the neutral arrow’s two heads increase line length. Familiar arrows can engage automatic and deliberate orienting. Describe this as predictive spatial cueing, not purely voluntary attention (Ristic and Kingstone, 2006, doi:10.1080/17470210500416367).')
h('Running a session', new_page=True)
h('Stimulus assignment', 2)
p('The enumeration-stimulus-pool-v2 pool contains 180 fixed coordinate configurations per numerosity. Pool versions are included in configuration IDs and exported rows. Keep experiment and pool versions identical across both sessions; v2 coordinates differ from v1. Each configuration occurs once per participant across both sessions. Forty-five assignment versions rotate nine signal delays and five cue-allocation slots; across those versions every configuration occurs at each delay three times valid, once invalid, and once neutral. Assign versions prospectively for recruitment balance. Seeded randomisation determines mask patterns, fixation durations, and trial order.')
h('Session procedure', 2)
p('A distinct Practice complete / Main experiment page shows trial/block counts, restates response rules, and requires Begin main experiment. Demonstration is labelled separately. All practice number pads share the calibrated main-trial position.')
for text in [
    'Open Start Experiment.command and use the local address it prints. Keep its window open. Libraries are bundled and participant records stay on the computer.',
    'The separate enumeration-jspsych folder retains the 20-trial demonstration. It uses numerosities 3 and 6, target exposures of 66.7, 100, 150, 250, and 400 ms, and the same single-pass training sequence. Base signal delays are 150 and 600 ms; short delays are increased where necessary so the target disappears by the tone. It is a preview rather than a balanced pilot dataset.',
    'Set the assumed MacBook display settings, then enter a participant code, session number, measured geometry, and assigned stimulus version. Confirm that the sound check is audible, then complete training and the nine blocks.',
    'Take self-paced breaks between blocks. Download the full JSON and trial CSV at completion. The setup page can export previous runs from browser recovery storage.'
]: p(text, 'List Number')
p('Main trials retain the fixed 66.7 ms target exposure and all nine original signal delays. Practice and demonstration vary exposure only for familiarisation and are excluded from main analyses. Their tones occur no earlier than the planned target offset; they can still overlap the mask.')
p('The updated HTML can also be opened directly for previews. Generated tones use jsPsych\'s audio context with its documented file-launch compatibility option. No audio file is fetched. Prefer the local server for a consistent research workflow; recovery storage can differ between launch methods.')

h('Data and inclusion rules', 2)
p('Full JSON retains configuration, metadata, scheduled trials, jsPsych records, display assumptions, cursor resets and positions, visual events, frame intervals, and scheduled audio timestamps. trial_start_method is automatic_cursor_reset; cursor_control is pointer_lock. Reset times are retained; cursor coordinates use viewport CSS pixels. CSV contains enumeration summaries. RT fields reference scheduled tone onset (rt and rt_tone_ms), array onset, and actual pad appearance (rt_pad_ms). response and reported_number contain the selected numeral, rather than a button index.')
p('correct_timely is numerical correctness for a timely first response and is null after timing failure. numerical_correct describes the initial report independently of timing. accuracy_all_scheduled is one only for a timely correct delivered trial and zero otherwise; undelivered scheduled trials are not observed errors. Report timing failures and response exclusions by cue, numerosity, and delay.')
p('Focus loss, page hiding, viewport resize, fullscreen exit, or interrupted audio ends and flags the affected trial. Interrupted trials are saved and not replayed. Recovery exports completed records rather than resuming an unfinished trial. Physical light and acoustic onset require equipment-based validation; browser timestamps are software measurements. Looking behaviour is not recorded, and performance curves alone do not uniquely establish an encoding mechanism.')

h('Research questions and model comparisons', new_page=True)
p('The primary question is whether spatial preparation changes the onset, rate, or eventual accuracy of enumeration, and whether that change depends on numerosity. A secondary question distinguishes benefits of valid cues relative to neutral from costs of invalid cues relative to neutral. Neutral indicates equal location probability and non-directional preparation; it is not an absence of attention.')
table(['Candidate account', 'Predicted performance pattern'], [
    ('Earlier onset', 'A cue effect on the point where performance begins rising above its baseline; a pure onset account does not require different rates or plateaus.'),
    ('Faster accrual', 'A cue effect on how rapidly performance approaches its plateau; a pure rate account does not require different onset or plateau parameters.'),
    ('Different plateau', 'A cue effect that persists at long processing times; a pure plateau account does not require different onset or rate parameters.')
])
p('Fit accuracy as a function of processing time using a suitable bounded likelihood and participant-level structure. Compare shared-parameter models, onset-only, rate-only, plateau-only, and combined accounts. Onset, rate, and plateau are estimated psychological quantities, not fixed experiment settings. A reliable separation requires parameter-recovery and model-recovery simulations using the planned allocation and plausible response variability.')
p('Plot accuracy against the nine scheduled delays, and retain actual array-to-mouse-press time for modelling and sensitivity analyses. Automatic cursor reset standardises position before fixation timing; movement afterwards is allowed and the observed position at pad onset is retained. Mouse movement to different buttons changes RT and can change available processing time despite identical scheduled delays. The 3000 ms deadline is for usability piloting and permits additional processing after the signal. Keep immediate-response instructions, evaluate tone-to-pad appearance delay, and validate an analysis of actual response times before confirmatory use; scheduled-delay curves alone cannot be treated as a tightly constrained response-signal procedure. Do not assume a universal 1/9 guessing baseline: 0 is available, and numerical response preferences can vary. Establish the baseline and fitting approach before confirmatory analysis.')
h('Evidence and decisions before the main study', 2)
p('Pilot arrow legibility, mask strength, mouse timing compliance, plateau coverage, and training burden. Freeze inclusion rules, the model family, numerosity contrasts, and the participant sample size after model recovery and precision evaluation. The current trial count is an implemented pilot allocation rather than an established power calculation.')
p('Methodological references: Giordano, McElree, and Carrasco (2009), Journal of Vision, 9(3), 30, doi:10.1167/9.3.30; Tuncok, Carrasco, and Winawer (2025), Nature Communications, 16, 8746, doi:10.1038/s41467-025-63795-3. Implementation references: jspsych.org/v8/overview/timeline/, jspsych.org/v8/developers/plugin-development/, and jspsych.org/v8/overview/running-experiments/.')

output = ROOT / 'docs/Experimental_design_and_parameters.docx'
doc.save(output)
print(output)
